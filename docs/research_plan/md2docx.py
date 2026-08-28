"""Convert 研究计划.md to a formatted .docx (simple Markdown subset)."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def set_cjk_font(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        h = doc.styles[name]
        h.font.name = "Times New Roman"
        h.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "黑体")


def add_runs(paragraph, text: str) -> None:
    pos = 0
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos : m.start()])
        paragraph.add_run(m.group(1)).bold = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def flush_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            p = table.cell(i, j).paragraphs[0]
            add_runs(p, cell)
            if i == 0:
                for r in p.runs:
                    r.bold = True
    doc.add_paragraph()


def convert(src: Path, dst: Path) -> None:
    doc = Document()
    set_cjk_font(doc)
    table_rows: list[list[str]] = []
    for raw in src.read_text(encoding="utf-8").splitlines():
        line = raw.rstrip()
        if line.startswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            if all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                continue
            table_rows.append(cells)
            continue
        flush_table(doc, table_rows)
        table_rows = []
        if not line.strip():
            continue
        if line == "---":
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            h = doc.add_heading(line[2:], level=1)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif re.match(r"^\d+\.\s", line.strip()):
            p = doc.add_paragraph(style="List Number")
            add_runs(p, re.sub(r"^\d+\.\s+", "", line.strip()))
        elif line.strip().startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, line.strip()[2:])
        else:
            p = doc.add_paragraph()
            add_runs(p, line.strip())
    flush_table(doc, table_rows)
    doc.save(dst)
    print(f"saved: {dst}")


if __name__ == "__main__":
    base = Path(__file__).parent
    convert(base / "研究计划.md", Path(sys.argv[1]) if len(sys.argv) > 1 else base / "研究计划.docx")
